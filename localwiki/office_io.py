"""Thin office document helpers for skills / exports (R3.11 / R3.17)."""

from __future__ import annotations

import io
import os
from pathlib import Path
from typing import Any

_PDF_MAX_CHARS = 100_000


def read_pdf(path: Path | str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def _find_cjk_font() -> Path | None:
    """Prefer common Windows / Noto CJK fonts for reportlab."""
    windir = os.environ.get("WINDIR") or r"C:\Windows"
    fonts = Path(windir) / "Fonts"
    candidates = [
        fonts / "msyh.ttc",
        fonts / "msyh.ttf",
        fonts / "msyhbd.ttc",
        fonts / "simhei.ttf",
        fonts / "simsun.ttc",
        fonts / "SIMSUN.TTC",
        fonts / "NotoSansSC-Regular.otf",
        fonts / "NotoSansCJKsc-Regular.otf",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for p in candidates:
        try:
            if p.is_file():
                return p
        except OSError:
            continue
    return None


def write_pdf(
    *,
    title: str = "",
    paragraphs: list[str] | None = None,
) -> bytes:
    """Write a simple multi-page PDF (R3.17). Raises ImportError if reportlab missing."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ImportError(
            "reportlab not installed; pip install reportlab into work venv"
        ) from exc

    paras = [str(p).strip() for p in (paragraphs or []) if str(p).strip()]
    title_s = (title or "").strip()
    blob = "\n\n".join(([title_s] if title_s else []) + paras)
    if len(blob) > _PDF_MAX_CHARS:
        paras = []
        remain = _PDF_MAX_CHARS - (len(title_s) + 2 if title_s else 0)
        for p in (paragraphs or []):
            s = str(p).strip()
            if not s:
                continue
            if remain <= 0:
                break
            if len(s) > remain:
                paras.append(s[:remain] + "…")
                break
            paras.append(s)
            remain -= len(s) + 2

    font_name = "Helvetica"
    font_path = _find_cjk_font()
    if font_path is not None:
        try:
            # TTC: subfontIndex 0 is usually the regular face
            pdfmetrics.registerFont(TTFont("LWCjk", str(font_path), subfontIndex=0))
            font_name = "LWCjk"
        except Exception:
            try:
                pdfmetrics.registerFont(TTFont("LWCjk", str(font_path)))
                font_name = "LWCjk"
            except Exception:
                font_name = "Helvetica"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LWTitle",
        parent=styles["Heading1"],
        fontName=font_name,
        fontSize=16,
        leading=22,
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "LWBody",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=11,
        leading=16,
        spaceAfter=8,
    )

    def _esc(text: str) -> str:
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )

    story: list[Any] = []
    if title_s:
        story.append(Paragraph(_esc(title_s), title_style))
        story.append(Spacer(1, 4 * mm))
    for p in paras:
        story.append(Paragraph(_esc(p), body_style))
    if not story:
        story.append(Paragraph(" ", body_style))
    doc.build(story)
    data = buf.getvalue()
    if not data.startswith(b"%PDF"):
        raise RuntimeError("reportlab produced non-PDF output")
    return data


def read_docx(path: Path | str) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip()


def write_docx(
    *,
    title: str = "",
    paragraphs: list[str] | None = None,
    bullet_points: list[str] | None = None,
) -> bytes:
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for p in paragraphs or []:
        if p:
            doc.add_paragraph(str(p))
    for b in bullet_points or []:
        if b:
            doc.add_paragraph(str(b), style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def read_xlsx(path: Path | str) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c) for c in row])
        return rows
    finally:
        wb.close()


def write_xlsx(rows: list[list[Any]], *, sheet_name: str = "Sheet1") -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Sheet1"
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_pptx(path: Path | str) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise ImportError("python-pptx not installed") from exc
    prs = Presentation(str(path))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
    return "\n".join(chunks).strip()


def pptx_available() -> bool:
    try:
        import pptx  # noqa: F401

        return True
    except ImportError:
        return False

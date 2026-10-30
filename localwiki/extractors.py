"""File extractors: MD/TXT/PDF/DOCX/XLSX (R2 / Personal Core)."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Literal

FormatKind = Literal["md", "txt", "pdf", "docx", "xlsx", "unknown"]

DEFAULT_MAX_BYTES = 100 * 1024 * 1024  # 100 MiB


class ExtractError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Locator:
    kind: str  # line | page | paragraph
    start: int | None = None
    end: int | None = None
    page: int | None = None
    paragraph_id: str | None = None
    excerpt: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {k: v for k, v in asdict(self).items() if v is not None and v != ""}


@dataclass
class ExtractResult:
    format: FormatKind
    text: str
    locators: list[Locator] = field(default_factory=list)
    sha256: str = ""
    title: str = ""
    meta: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "format": self.format,
            "text": self.text,
            "locators": [L.to_dict() for L in self.locators],
            "sha256": self.sha256,
            "title": self.title,
            "meta": self.meta,
        }


def detect_format(filename: str, data: bytes | None = None) -> FormatKind:
    name = filename.lower()
    if name.endswith((".md", ".markdown")):
        return "md"
    if name.endswith(".txt"):
        return "txt"
    if name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if data and data[:4] == b"%PDF":
        return "pdf"
    if data and data[:2] == b"PK":
        head = data[:8192]
        if b"word/" in head:
            return "docx"
        if b"xl/" in head or b"xl\\" in head:
            return "xlsx"
    return "unknown"


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_bytes(
    data: bytes,
    *,
    filename: str = "upload.bin",
    max_bytes: int = DEFAULT_MAX_BYTES,
) -> ExtractResult:
    if len(data) > max_bytes:
        raise ExtractError("E_INVALID_INPUT", f"file exceeds {max_bytes} bytes")
    digest = hashlib.sha256(data).hexdigest()
    kind = detect_format(filename, data)
    if kind == "unknown":
        raise ExtractError("E_INVALID_INPUT", f"unsupported format: {filename}")

    if kind in ("md", "txt"):
        text = _decode_text(data)
        locs: list[Locator] = []
        lines = text.splitlines()
        for i, line in enumerate(lines, start=1):
            if line.strip():
                locs.append(Locator(kind="line", start=i, end=i, excerpt=line[:200]))
        title = _title_from_text(text, filename)
        return ExtractResult(
            format=kind,
            text=text,
            locators=locs[:500],
            sha256=digest,
            title=title,
            meta={"encoding_guess": "auto"},
        )

    if kind == "pdf":
        return _extract_pdf(data, digest=digest, filename=filename)

    if kind == "docx":
        return _extract_docx(data, digest=digest, filename=filename)

    if kind == "xlsx":
        return _extract_xlsx(data, digest=digest, filename=filename)

    raise ExtractError("E_INVALID_INPUT", f"unsupported format: {kind}")


def extract_path(path: Path | str, *, max_bytes: int = DEFAULT_MAX_BYTES) -> ExtractResult:
    p = Path(path)
    data = p.read_bytes()
    return extract_bytes(data, filename=p.name, max_bytes=max_bytes)


def _title_from_text(text: str, filename: str) -> str:
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("# "):
            return s[2:].strip()
        if s:
            return s[:80]
    return Path(filename).stem


def _extract_pdf(data: bytes, *, digest: str, filename: str) -> ExtractResult:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractError("E_DEP_MISSING", "pypdf required") from exc
    import io

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    locs: list[Locator] = []
    for i, page in enumerate(reader.pages, start=1):
        t = page.extract_text() or ""
        parts.append(t)
        if t.strip():
            locs.append(
                Locator(kind="page", page=i, excerpt=t.strip()[:200])
            )
    text = "\n\n".join(parts)
    return ExtractResult(
        format="pdf",
        text=text,
        locators=locs,
        sha256=digest,
        title=_title_from_text(text, filename),
        meta={"pages": len(reader.pages)},
    )


def _extract_docx(data: bytes, *, digest: str, filename: str) -> ExtractResult:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractError("E_DEP_MISSING", "python-docx required") from exc
    import io

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    locs: list[Locator] = []
    for idx, para in enumerate(doc.paragraphs):
        t = para.text or ""
        if not t.strip():
            continue
        pid = f"p{idx}"
        parts.append(t)
        locs.append(Locator(kind="paragraph", paragraph_id=pid, excerpt=t[:200]))
    text = "\n\n".join(parts)
    return ExtractResult(
        format="docx",
        text=text,
        locators=locs,
        sha256=digest,
        title=_title_from_text(text, filename),
        meta={"paragraphs": len(locs)},
    )


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_xlsx(data: bytes, *, digest: str, filename: str) -> ExtractResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ExtractError("E_DEP_MISSING", "openpyxl required") from exc
    import io

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractError("E_INVALID_INPUT", f"cannot read xlsx: {exc}") from exc

    parts: list[str] = []
    locs: list[Locator] = []
    sheet_count = 0
    row_count = 0
    for sheet in wb.worksheets:
        sheet_count += 1
        sheet_lines: list[str] = []
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [_cell_str(c) for c in row]
            if not any(cells):
                continue
            line = "\t".join(cells)
            sheet_lines.append(line)
            row_count += 1
            if len(locs) < 500:
                locs.append(
                    Locator(
                        kind="line",
                        start=r_idx,
                        end=r_idx,
                        excerpt=f"[{sheet.title}] {line[:180]}",
                    )
                )
        if sheet_lines:
            parts.append(f"## {sheet.title}\n" + "\n".join(sheet_lines))
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass

    text = "\n\n".join(parts).strip()
    if not text:
        text = f"(empty spreadsheet: {Path(filename).name})"
    return ExtractResult(
        format="xlsx",
        text=text,
        locators=locs,
        sha256=digest,
        title=_title_from_text(text, filename),
        meta={"sheets": sheet_count, "rows": row_count},
    )


_SLUG_RE = re.compile(r"[^\w\u4e00-\u9fff]+", re.UNICODE)


def slugify(title: str, *, fallback: str = "page") -> str:
    s = _SLUG_RE.sub("-", title.strip()).strip("-").lower()
    return (s[:60] or fallback)


def write_locators_json(path: Path, locators: list[Locator]) -> None:
    path.write_text(
        json.dumps([L.to_dict() for L in locators], ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
